"""Arma los anexos E, F, I y J a partir de los documentos existentes (sin tocar los originales).

Uso: python scripts/build_anexos_existentes.py   (requiere python-docx y docxcompose)
Salida: ../Anexos/
"""
import copy
import re
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

BASE = Path(__file__).resolve().parents[2]      # carpeta "Trabajo de Grado (1)"
OUT = BASE / "Anexos"
OUT.mkdir(exist_ok=True)
CAP = re.compile(r"\b(Tabla|Figura) (\d+)\b")


def clone_before(anchor, template, text, bold=None):
    """Inserta antes de `anchor` un párrafo con el formato de `template` y el texto dado."""
    new = copy.deepcopy(template._p)
    anchor._p.addprevious(new)
    para = docx.text.paragraph.Paragraph(new, anchor._parent)
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)
    para.runs[0].text = text
    if bold is not None:
        para.runs[0].bold = bold
    return para


def set_text(para, old, new):
    for r in para.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    return False


def replace_everywhere(d, old, new):
    hits = 0
    for para in iter_paragraphs(d):
        hits += set_text(para, old, new)
    assert hits, f"no encontrado: {old!r}"


def iter_paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def renumber(d, letter, offset=0):
    """Tabla N / Figura N -> Tabla X.(N+offset). Devuelve el máximo N por tipo."""
    mx = {"Tabla": 0, "Figura": 0}
    for para in iter_paragraphs(d):
        for r in para.runs:
            def sub(m):
                k, v = m.group(1), int(m.group(2))
                mx[k] = max(mx[k], v)
                return f"{k} {letter}.{v + offset.get(k, 0) if isinstance(offset, dict) else v + offset}"
            new = CAP.sub(sub, r.text)
            if new != r.text:          # asignar solo si cambia: run.text= borra imágenes del run
                r.text = new
    left = [p.text for p in iter_paragraphs(d) if re.search(r"\b(Tabla|Figura) \d", p.text)]
    assert not left, f"quedaron referencias sin renumerar (partidas entre runs): {left[:3]}"
    return mx


def black_headings(d):
    """Los títulos heredan el color del estilo Heading de Word (azul); NTC 1486 los usa en negro."""
    from docx.shared import RGBColor
    for para in d.paragraphs:
        st = para._p.pPr.pStyle.val if (para._p.pPr is not None and para._p.pPr.pStyle is not None) else ""
        if st.lower().startswith("heading") or st.lower().startswith("ttulo"):
            for r in para.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)


def add_title(d, lines):
    first = d.paragraphs[0]
    for i, txt in enumerate(lines):
        clone_before(first, first, txt, bold=True)


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def merge(master, other, path):
    black_headings(master); black_headings(other)
    page_break(master)
    comp = Composer(master)
    comp.append(other)
    comp.save(path)
    print("->", path.name)


# --- Anexo E: Proceso OE1 + Validación externa -------------------------------
e1 = docx.Document(BASE / "Proceso_Objetivo1.docx")
e2 = docx.Document(BASE / "Validacion_Externa_pyscan.docx")
for para in e1.paragraphs:
    if para.text.startswith("Se obtuvo un dataset de 4.000 muestras"):
        para.add_run(" Nota: este fue el dataset de la primera versión del modelo; posteriormente se amplió "
                     "con una muestra aleatoria de benignos de PyPI (composición final en el Anexo G).")
        break
else:
    raise SystemExit("párrafo del dataset no encontrado en Proceso_Objetivo1")
for para in e1.paragraphs:
    if "Se registró como línea de mejora" in para.text:
        para.add_run(" Esa mejora se aplicó después: con la muestra aleatoria de benignos la dependencia del nombre "
                     "quedó acotada (Anexo H, estudio de ablación).")
        break
else:
    raise SystemExit("párrafo de línea de mejora no encontrado")
replace_everywhere(e2, "Recall 0,97 · F1 0,975 · FP 2,3 %", "Recall 0,945 · F1 0,933 · FP 6,1 %")
mx = renumber(e1, "E")
renumber(e2, "E", offset=mx)
add_title(e1, ["ANEXO E", "PROCESO DE IDENTIFICACIÓN DE VECTORES Y VALIDACIÓN EXTERNA DE LAS SEÑALES"])
merge(e1, e2, OUT / "Anexo_E_Proceso_OE1_y_validacion_externa.docx")

# --- Anexo F: Proceso OE2 (ADR) + Comparativo de arquitecturas ---------------
f1 = docx.Document(BASE / "Proceso_Objetivo2.docx")
f2 = docx.Document(BASE / "Comparativo_Arquitecturas.docx")
mx = renumber(f1, "F")
renumber(f2, "F", offset=mx)
add_title(f1, ["ANEXO F", "DECISIONES DE DISEÑO Y COMPARATIVO DE ARQUITECTURAS"])
merge(f1, f2, OUT / "Anexo_F_Decisiones_de_arquitectura.docx")

# --- Anexo I: Manual técnico --------------------------------------------------
mt = docx.Document(BASE / "Manual_Tecnico.docx")
replace_everywhere(mt, "La suite consta de 58 pruebas", "La suite consta de 64 pruebas")
replace_everywhere(mt, "(anti Zip-Slip, symlink y zip-bomb)",
                   "(anti Zip-Slip y zip-bomb; omite los enlaces simbólicos y duros sin escribirlos en disco)")
replace_everywhere(mt, "Orquesta el pipeline y expone los comandos.",
                   "Orquesta el pipeline y expone los comandos (scan, check-name, info, precommit, version).")
for para in mt.paragraphs:
    if para.text.startswith("El clasificador es supervisado"):
        para.add_run(" Tras la validación se seleccionó Random Forest con umbral 0,45 "
                     "(Recall 0,945 y F1 0,933 en el hold-out; detalle en el Anexo H).")
    if para.text.startswith("python scripts/collect_benign.py --limit 2000"):
        tpl = para
nxt = tpl._p.getnext()
anchor = docx.text.paragraph.Paragraph(nxt, tpl._parent)
for line in ["python scripts/fetch_random_pypi.py --limit 1000 --seed 42",
             "python scripts/collect_benign.py --names-file data/random_pypi_names.txt --out data/benign/pypi_random --limit 1000"]:
    clone_before(anchor, tpl, line)
renumber(mt, "I")
add_title(mt, ["ANEXO I"])
black_headings(mt)
mt.save(OUT / "Anexo_I_Manual_tecnico.docx"); print("-> Anexo_I_Manual_tecnico.docx")

# --- Anexo J: Manual de usuario -----------------------------------------------
mu = docx.Document(BASE / "Manual_Usuario.docx")
P = mu.paragraphs
h2_tpl = next(p for p in P if p.text.startswith("4.4 Otros comandos"))
code_tpl = next(p for p in P if p.text.startswith("python -m pyscan.cli check-name"))
body_tpl = next(p for p in P if p.text.startswith("Lee todas las dependencias"))
anchor = next(p for p in P if p.text.startswith("5. Cómo interpretar"))
clone_before(anchor, h2_tpl, "4.5 Revisar las dependencias antes de cada commit (pre-commit / CI)")
clone_before(anchor, code_tpl, "python -m pyscan.cli precommit requirements.txt")
clone_before(anchor, code_tpl, "python -m pyscan.cli precommit --allow-unscanned requirements.txt")
clone_before(anchor, body_tpl, "Analiza las dependencias del archivo y bloquea el commit (código de salida 1) si alguna resulta "
             "MALICIOSA, si su nombre es sospechoso de typosquatting o si no se pudo analizar. Con --allow-unscanned, "
             "las dependencias sin analizar solo se avisan. Para integrarlo con la herramienta pre-commit, el "
             "proyecto incluye el archivo .pre-commit-hooks.yaml (ver docs/PRECOMMIT.md).")
clone_before(anchor, h2_tpl, "4.6 Interfaz web local y ejecutable")
clone_before(anchor, code_tpl, "python scripts/webapp.py")
clone_before(anchor, body_tpl, "Abre una interfaz en el navegador (http://127.0.0.1:5000) que usa el mismo motor de análisis: "
             "se pegan los nombres de los paquetes o el contenido de un requirements.txt y se obtiene el veredicto "
             "de cada dependencia. Solo funciona en el equipo local y nunca ejecuta el código de los paquetes. "
             "También puede distribuirse como ejecutable de doble clic (build_exe.bat en Windows; ver docs/EJECUTABLE.md).")
replace_everywhere(mu, "6. Ejemplo de salida", "6. Ejemplo de salida (ilustrativo)")
trouble = next(p for p in mu.paragraphs if p.text.startswith("Un paquete no se pudo analizar"))
after = trouble._p.getnext()
clone_before(docx.text.paragraph.Paragraph(after, trouble._parent), trouble,
             "El hook de pre-commit bloquea con «NO se pudieron analizar»: revise la conexión o el paquete indicado; "
             "si decide continuar, use --allow-unscanned.", bold=False)
renumber(mu, "J")
add_title(mu, ["ANEXO J"])
black_headings(mu)
mu.save(OUT / "Anexo_J_Manual_de_usuario.docx"); print("-> Anexo_J_Manual_de_usuario.docx")
