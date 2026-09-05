#!/usr/bin/env python3
"""Genera las diapositivas de la ponencia sobre la plantilla institucional."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from copy import deepcopy

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x5C, 0x9E)
ORANGE = RGBColor(0xE8, 0x79, 0x1E)
GRAY = RGBColor(0x3A, 0x3A, 0x3A)
LGRAY = RGBColor(0x5B, 0x5B, 0x5B)
CARD = RGBColor(0xEE, 0xF1, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation("plantilla.pptx")
LAYOUT = prs.slides[1].slide_layout  # "Título y objetos" con el marco institucional
COVER_LAYOUT = prs.slides[0].slide_layout

CX = Inches(0.85)          # margen izquierdo del área de contenido
CW = Inches(8.45)          # ancho del área de contenido
TITLE_Y = Inches(1.5)
BODY_Y = Inches(2.35)
BODY_H = Inches(4.0)


def _clear(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def add_slide():
    s = prs.slides.add_slide(LAYOUT)
    _clear(s)
    return s


def set_font(run, size, color=GRAY, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.name = font; run.font.color.rgb = color


def title(slide, text, color=NAVY, size=26):
    tb = slide.shapes.add_textbox(CX, TITLE_Y, CW, Inches(0.8))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    set_font(r, size, color, bold=True)
    return tb


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.05); tf.margin_bottom = Inches(0.05)
    return tb, tf


def bullets(slide, items, x=CX, y=BODY_Y, w=CW, h=BODY_H, size=16, gap=10):
    tb, tf = textbox(slide, x, y, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        p.level = lvl
        r = p.add_run(); r.text = ("•  " if lvl == 0 else "–  ") + it
        set_font(r, size if lvl == 0 else size - 1, GRAY if lvl == 0 else LGRAY)
    return tb


def card(slide, x, y, w, h, fill=CARD, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line: sp.line.color.rgb = line; sp.line.width = Pt(1)
    else: sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def stat_card(slide, x, y, w, num, label, numcolor=NAVY):
    h = Inches(1.7)
    card(slide, x, y, w, h)
    tb, tf = textbox(slide, x, y + Inches(0.15), w, Inches(0.9), anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; set_font(r, 44, numcolor, bold=True)
    tb2, tf2 = textbox(slide, x + Inches(0.1), y + Inches(1.02), w - Inches(0.2), Inches(0.6),
                       anchor=MSO_ANCHOR.TOP)
    p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label; set_font(r2, 12.5, GRAY, bold=True)


# ============ PORTADA (slide 1 existente: logo) → dejar tal cual ============
# ============ SLIDE 2: TÍTULO (reusa slide2, limpiándola) ============
s = prs.slides[1]
_clear(s)
tb, tf = textbox(s, CX, Inches(1.7), CW, Inches(1.9))
p = tf.paragraphs[0]; r = p.add_run()
r.text = "Detección de paquetes maliciosos en PyPI: un mapeo estructurado de la evidencia, las prácticas de evaluación y los vacíos abiertos (2020–2026)"
set_font(r, 24, NAVY, bold=True)
tb2, tf2 = textbox(s, CX, Inches(3.9), CW, Inches(1.6))
autores = [
    ("Andrés Felipe Sanguino Cubillos · Juan David Gutiérrez Reyes", 15, NAVY, True),
    ("Solangie Paola Garavito Mendivelso", 15, NAVY, True),
    ("Programa de Ingeniería de Sistemas y Computación", 13, GRAY, False),
    ("Universidad Católica de Colombia — Bogotá D.C.", 13, GRAY, False),
    ("Modalidad: Ponencia   |   Línea: Gestión y Tecnología al Servicio de la Sociedad   |   ODS 9", 12, LGRAY, False),
]
for i, (txt, sz, col, bd) in enumerate(autores):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    p.space_after = Pt(6); r = p.add_run(); r.text = txt; set_font(r, sz, col, bold=bd)

# ============ SLIDE 3: CONTEXTO Y PROBLEMA ============
s = add_slide(); title(s, "Contexto: la confianza en los repositorios de paquetes")
bullets(s, [
    "El desarrollo moderno depende de repositorios públicos y de la resolución automática de dependencias.",
    "Ese modelo traslada la confianza a los nombres de paquetes, las cuentas de mantenedores y los procedimientos que se ejecutan al instalar.",
    "Los atacantes explotan esa confianza: suplantación de nombres, confusión de dependencias, compromiso de cuentas, actualizaciones maliciosas y cargas ofuscadas.",
    "PyPI es un caso crítico: un paquete puede ejecutar código en la construcción, la instalación, la importación o la ejecución.",
])

# ============ SLIDE 4: EL PROBLEMA DE FONDO ============
s = add_slide(); title(s, "El problema de fondo")
bullets(s, [
    "El comportamiento malicioso se distribuye entre metadatos, archivos de instalación, módulos, recursos codificados y descargas externas.",
    "Las etiquetas a nivel de paquete ocultan un problema de localización de la lógica dañina.",
    "Un incidente reciente mostró que un adversario paciente puede combinar ingeniería social, un historial legítimo de contribuciones y lógica de construcción oculta para eludir los controles perimetrales.",
], size=16.5, gap=14)

# ============ SLIDE 5: MOTIVACIÓN / BRECHA ============
s = add_slide(); title(s, "La brecha: desempeño alto, pero no comparable")
bullets(s, [
    "La literatura responde con múltiples enfoques: reglas, ML supervisado, representaciones aprendidas, análisis de programa, ejecución dinámica y sistemas basados en LLM.",
    "El desempeño reportado suele ser alto, pero ordenar las propuestas es inseguro:",
    ("colecciones de malware distintas, muestreo de benignos distinto, unidades de análisis diferentes,", 1),
    ("balances de clase y particiones diferentes; algunos evalúan un banco histórico y otros paquetes hallados en despliegue.", 1),
    "Los estudios secundarios existentes no ofrecen una síntesis actual, centrada en PyPI, sobre la evidencia de los detectores y la comparabilidad.",
])

# ============ SLIDE 6: OBJETIVO Y PREGUNTAS ============
s = add_slide(); title(s, "Objetivo y preguntas de investigación")
tb, tf = textbox(s, CX, BODY_Y, CW, Inches(0.9))
p = tf.paragraphs[0]; r = p.add_run()
r.text = "Sintetizar la evidencia sobre detección de paquetes maliciosos en PyPI mediante un mapeo estructurado (2020–2026)."
set_font(r, 16, GRAY, italic=True)
qs = [("RQ1", "¿Qué paradigmas de detección predominan?"),
      ("RQ2", "¿Qué señales y unidades de análisis emplean?"),
      ("RQ3", "¿Qué vacíos de evidencia y operacionales impiden la comparación justa y el despliegue?")]
y = Inches(3.05)
for tag, q in qs:
    card(s, CX, y, Inches(1.1), Inches(0.62), fill=NAVY)
    tbx, tfx = textbox(s, CX, y, Inches(1.1), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
    pp = tfx.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = tag; set_font(rr, 18, WHITE, bold=True)
    tb2, tf2 = textbox(s, CX + Inches(1.3), y, CW - Inches(1.3), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
    p2 = tf2.paragraphs[0]; r2 = p2.add_run(); r2.text = q; set_font(r2, 16, GRAY)
    y += Inches(0.74)

# ============ SLIDE 7: MÉTODO ============
s = add_slide(); title(s, "Método: mapeo estructurado")
bullets(s, [
    "Mapeo estructurado (no revisión sistemática exhaustiva): fuentes acotadas, se incluyen preprints, sin doble revisión ciega.",
    "Alcance definido con PICOC (población, intervención, comparación, resultados, contexto).",
    "Corte editorial: julio de 2026 · publicaciones de 2020 a 2026.",
    "Búsqueda en 2 índices de citación + 5 bases de editorial/repositorio, con consulta canónica por base.",
    "Complementada con rastreo hacia atrás sobre una taxonomía de ataques y una colección conocida de paquetes maliciosos.",
    "Síntesis narrativa: la heterogeneidad de datos y protocolos hace inapropiado un metaanálisis.",
], size=15, gap=8)

# ============ SLIDE 8: CORPUS ============
s = add_slide(); title(s, "Corpus analizado")
w = Inches(2.55); gap = Inches(0.35); x0 = CX + Inches(0.35)
stat_card(s, x0, Inches(2.5), w, "29", "trabajos académicos")
stat_card(s, x0 + w + gap, Inches(2.5), w, "23", "propuestas de detección", numcolor=ORANGE)
stat_card(s, x0 + 2*(w + gap), Inches(2.5), w, "6", "caracterización / datasets")
bullets(s, [
    "Dos herramientas operativas se conservaron solo como líneas base (no se cuentan como estudios).",
    "Por cada trabajo se registró: propósito, paradigma, señal, unidad de análisis, procedencia y balance de datos, partición, métricas, robustez, costo y disponibilidad de artefactos.",
], y=Inches(4.55), size=14.5, gap=8)

# ============ SLIDE 9: RESULTADO 1 — PARADIGMAS ============
s = add_slide(); title(s, "RQ1 · Paradigmas de detección")
w = Inches(2.55); gap = Inches(0.35); x0 = CX + Inches(0.35)
stat_card(s, x0, Inches(2.4), w, "13", "ML clásico y representaciones aprendidas")
stat_card(s, x0 + w + gap, Inches(2.4), w, "7", "LLM, RAG y agentes", numcolor=ORANGE)
stat_card(s, x0 + 2*(w + gap), Inches(2.4), w, "3", "análisis de programa / integridad")
bullets(s, [
    "No hay polarización estricta entre heurísticas y LLM: la evidencia describe un continuo.",
    "Reglas económicas e interpretables; representaciones que mejoran la generalización; análisis de programa que mejora la trazabilidad; LLM que amplían el contexto con más dependencias.",
    "Ningún paradigma domina bajo todas las restricciones.",
], y=Inches(4.45), size=14.5, gap=7)

# ============ SLIDE 10: RESULTADO 2 — SEÑALES Y UNIDADES ============
s = add_slide(); title(s, "RQ2 · Señales y unidades de análisis")
bullets(s, [
    "Cinco familias de señales:",
    ("metadatos del repositorio · señales léxicas y estructurales del código · representaciones conductuales · evidencia dinámica en entornos aislados · representaciones semánticas.", 1),
    "La unidad de análisis es tan determinante como la señal: paquete, publicación, archivo, archivo de instalación, módulo, secuencia de llamadas y sentencia.",
    "Las predicciones a nivel de archivo no son comparables con métricas a nivel de paquete sin una regla de agregación declarada.",
    "Riesgo de fuga de datos entre versiones de un mismo paquete.",
], size=15.5, gap=9)

# ============ SLIDE 11: RESULTADO 3 — VACÍO DOMINANTE ============
s = add_slide(); title(s, "RQ3 · El vacío dominante: la comparabilidad")
card(s, CX, Inches(2.4), CW, Inches(1.0), fill=CARD)
tbx, tfx = textbox(s, CX + Inches(0.2), Inches(2.4), CW - Inches(0.4), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
pp = tfx.paragraphs[0]; rr = pp.add_run()
rr.text = "La procedencia de los datos, las unidades, el balance de clases, la calidad de las etiquetas y las particiones se reportan de forma inconsistente."
set_font(rr, 16, NAVY, bold=True)
bullets(s, [
    "La validación temporal (entrenar en el pasado, evaluar en el futuro) es poco frecuente.",
    "En conjunto, estos vacíos impiden comparar de forma justa y trasladar los resultados al despliegue.",
], y=Inches(3.7), size=16, gap=10)

# ============ SLIDE 12: CONTRIBUCIONES ============
s = add_slide(); title(s, "Contribuciones")
items = [
    ("1", "Mapa de evidencia corregido", "separa estudios académicos de herramientas operativas, y el propósito del estudio de la técnica de detección."),
    ("2", "Taxonomía de señales y unidades", "desde los metadatos hasta el contexto semántico."),
    ("3", "Análisis de incomparabilidad", "por qué las métricas publicadas no son directamente comparables."),
    ("4", "Protocolo mínimo de reporte", "requisitos comunes para evaluar futuros detectores."),
]
y = Inches(2.2)
for num, h, d in items:
    card(s, CX, y, Inches(0.7), Inches(0.66), fill=ORANGE)
    tbx, tfx = textbox(s, CX, y, Inches(0.7), Inches(0.66), anchor=MSO_ANCHOR.MIDDLE)
    pp = tfx.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = num; set_font(rr, 20, WHITE, bold=True)
    tb2, tf2 = textbox(s, CX + Inches(0.9), y - Inches(0.02), CW - Inches(0.9), Inches(0.85))
    p1 = tf2.paragraphs[0]; r1 = p1.add_run(); r1.text = h + ": "; set_font(r1, 15.5, NAVY, bold=True)
    r2 = p1.add_run(); r2.text = d; set_font(r2, 14.5, GRAY)
    y += Inches(0.8)

# ============ SLIDE 13: PROTOCOLO MÍNIMO ============
s = add_slide(); title(s, "Protocolo mínimo de reporte (8 dimensiones)")
dims = ["Identidad del conjunto de datos", "Unidad de análisis", "Partición y fuga",
        "Etiquetas y desbalance", "Métricas", "Robustez", "Eficiencia y costo",
        "Artefactos reproducibles"]
cols, cw2, ch2 = 2, Inches(4.05), Inches(0.6)
x0, y0, gx, gy = CX, Inches(2.3), Inches(0.35), Inches(0.16)
for i, d in enumerate(dims):
    r, c = divmod(i, cols)
    x = x0 + c * (cw2 + gx); y = y0 + r * (ch2 + gy)
    card(s, x, y, cw2, ch2, fill=CARD)
    tbx, tfx = textbox(s, x + Inches(0.15), y, cw2 - Inches(0.2), ch2, anchor=MSO_ANCHOR.MIDDLE)
    pp = tfx.paragraphs[0]
    rn = pp.add_run(); rn.text = f"{i+1}. "; set_font(rn, 14, ORANGE, bold=True)
    rr = pp.add_run(); rr.text = d; set_font(rr, 14, GRAY, bold=True)
tb, tf = textbox(s, Inches(2.35), Inches(5.35), Inches(6.9), Inches(0.5))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT; r = p.add_run()
r.text = "Agnóstico al modelo: aplica desde las reglas hasta los sistemas basados en LLM."
set_font(r, 13, LGRAY, italic=True)

# ============ SLIDE 14: CONCLUSIONES Y TRABAJO FUTURO ============
s = add_slide(); title(s, "Conclusiones y trabajo futuro")
bullets(s, [
    "El campo es un continuo de paradigmas; ninguno domina bajo todas las restricciones.",
    "El principal obstáculo no es la técnica, sino la comparabilidad de las evaluaciones.",
    "El protocolo mínimo ofrece una base concreta para evaluar detectores en PyPI e interpretar las afirmaciones de desempeño sin ordenar estudios incompatibles.",
    "Trabajo futuro: implementar un banco de pruebas compartido bajo el protocolo y evaluar detectores escalonados usando paquetes históricos y monitoreo prospectivo del repositorio.",
], size=16, gap=12)

# ============ SLIDE 15: CIERRE ============
s = add_slide()
tb, tf = textbox(s, CX, Inches(2.6), CW, Inches(1.2), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "¡Gracias!"; set_font(r, 40, NAVY, bold=True)
tb2, tf2 = textbox(s, CX, Inches(3.9), CW, Inches(1.2), anchor=MSO_ANCHOR.TOP)
for i, txt in enumerate(["¿Preguntas?", "afsanguino15@ucatolica.edu.co · jdgutierrez017@ucatolica.edu.co",
                         "spgaravito@ucatolica.edu.co"]):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(6)
    r = p.add_run(); r.text = txt; set_font(r, 16 if i == 0 else 13, GRAY, bold=(i == 0))

prs.save("Ponencia_Deteccion_PyPI.pptx")
print("PPTX guardado con", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
